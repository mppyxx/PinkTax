"""Step 5 - Build the Excel workbook (excel/PinkTax_Analysis.xlsx).

Sheets
  About            - what each sheet contains
  Data             - the full cleaned master table as an Excel Table (tbl_Products)
  Category_Summary - COUNTIFS / AVERAGEIFS formulas computed live from tbl_Products,
                     plus medians and test results from Python, conditional formatting
  Segment_Pivot    - pivot-style gender x segment summary built with formulas
  Matched_Pairs    - brand-matched comparisons with a data bar on the price gap
  Stats_Tests      - regression and matched-pair test results
  Charts           - native Excel charts
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import EXCEL, PROCESSED, TABLES

PINK, BLUE, INK = "D6457F", "2A78D6", "1F1F1D"
HEAD = PatternFill("solid", fgColor="3B2A33")
HFONT = Font(bold=True, color="FFFFFF")
THIN = Border(bottom=Side(style="thin", color="E6E5E0"))
RUPEE = '"₹"#,##0'
PCT = '0.0"%"'


def header(ws, row, cols):
    for j, name in enumerate(cols, 1):
        c = ws.cell(row=row, column=j, value=name)
        c.fill, c.font = HEAD, HFONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def widths(ws, w):
    for i, v in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = v


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=15, color=INK)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="5C5B57")


def main():
    df = pd.read_csv(PROCESSED / "pinktax_master.csv", low_memory=False)
    cat = pd.read_csv(TABLES / "category_premium_mrp.csv")
    pairs = pd.read_csv(TABLES / "brand_matched_pairs.csv")
    reg = pd.read_csv(TABLES / "regression_fixed_effects.csv")
    psum = pd.read_csv(TABLES / "brand_matched_summary.csv")

    wb = Workbook()

    # ---------------------------------------------------------------- About
    ws = wb.active
    ws.title = "About"
    title(ws, "Pink Tax Analysis - Excel Workbook",
          "Gender-based price differences in Indian online retail (Myntra, BigBasket, Amazon.in)")
    rows = [
        ("Data", "Full cleaned dataset as an Excel Table named tbl_Products (filter, sort, pivot)."),
        ("Category_Summary", "Live COUNTIFS / AVERAGEIFS formulas per category + medians, CIs, p-values."),
        ("Segment_Pivot", "Pivot-style summary (segment x gender) built with formulas."),
        ("Matched_Pairs", "Same brand + same category comparisons of the men's and women's lines."),
        ("Stats_Tests", "Fixed-effects regression and matched-pair test results."),
        ("Charts", "Native Excel charts of the key results."),
        ("", ""),
        ("Premium %", "(women's price / men's price - 1) x 100. Positive = women pay more."),
        ("MRP", "Maximum Retail Price (list price printed on the product)."),
        ("Selling price", "Price after the platform discount."),
        ("Currency", "Indian Rupees (₹)."),
    ]
    for i, (k, v) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    widths(ws, [20, 100])

    # ---------------------------------------------------------------- Data
    ws = wb.create_sheet("Data")
    cols = ["source", "product_id", "product_name", "brand", "segment", "category", "gender",
            "age_group", "mrp", "selling_price", "discount_pct", "size_value", "size_unit",
            "unit_price_100", "rating", "rating_count"]
    d = df[cols].copy()
    d["product_name"] = d["product_name"].astype(str).str.slice(0, 120)
    d["discount_pct"] = d["discount_pct"].round(2)
    d["unit_price_100"] = d["unit_price_100"].round(2)
    ws.append(cols)
    for r in d.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])
    ref = f"A1:{get_column_letter(len(cols))}{len(d) + 1}"
    tbl = Table(displayName="tbl_Products", ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)
    ws.freeze_panes = "A2"
    widths(ws, [11, 14, 50, 18, 14, 18, 9, 10, 10, 12, 12, 10, 9, 13, 8, 12])
    for col in ("I", "J"):
        for c in ws[col][1:]:
            c.number_format = RUPEE
    n = len(d) + 1
    rng = lambda letter: f"Data!${letter}$2:${letter}${n}"
    SRC, CAT, GEN, AGE, MRP, SP, DISC = (rng(x) for x in "AFGHIJK")

    # ---------------------------------------------------------------- Category_Summary
    ws = wb.create_sheet("Category_Summary")
    title(ws, "Category-level comparison",
          "Columns E-J are live Excel formulas on the Data sheet; K-P come from the Python analysis.")
    hdr = ["Source", "Age group", "Segment", "Category", "Products (Men)", "Products (Women)",
           "Avg MRP (Men)", "Avg MRP (Women)", "Avg-price premium %", "Avg discount gap (pp)",
           "Median MRP (Men)", "Median MRP (Women)", "Median premium %", "95% CI low %",
           "95% CI high %", "p-value (BH adj.)", "Result"]
    header(ws, 4, hdr)
    c = cat.sort_values("premium_median_pct", ascending=False).reset_index(drop=True)
    for i, r in c.iterrows():
        row = 5 + i
        crit = f'{SRC},$A{row},{AGE},$B{row},{CAT},$D{row}'
        ws.append([
            r.source, r.age_group, r.segment, r.category,
            f'=COUNTIFS({crit},{GEN},"Men")', f'=COUNTIFS({crit},{GEN},"Women")',
            f'=AVERAGEIFS({MRP},{crit},{GEN},"Men")', f'=AVERAGEIFS({MRP},{crit},{GEN},"Women")',
            f"=(H{row}/G{row}-1)*100",
            f'=AVERAGEIFS({DISC},{crit},{GEN},"Women")-AVERAGEIFS({DISC},{crit},{GEN},"Men")',
            r.median_men, r.median_women, round(r.premium_median_pct, 2),
            round(r.ci_low_pct, 2), round(r.ci_high_pct, 2), float(f"{r.p_adj_bh:.3g}"), r.direction,
        ])
    last = 4 + len(c)
    for col, fmt in {"G": RUPEE, "H": RUPEE, "K": RUPEE, "L": RUPEE, "I": PCT, "J": '0.0',
                     "M": PCT, "N": PCT, "O": PCT, "P": "0.00E+00"}.items():
        for cell in ws[col][4:last]:
            cell.number_format = fmt
    for col in ("I", "M"):
        ws.conditional_formatting.add(f"{col}5:{col}{last}", ColorScaleRule(
            start_type="num", start_value=-40, start_color=BLUE,
            mid_type="num", mid_value=0, mid_color="F5F5F2",
            end_type="num", end_value=40, end_color=PINK))
    ws.conditional_formatting.add(f"Q5:Q{last}", CellIsRule(
        operator="equal", formula=['"Women pay more"'], font=Font(bold=True, color=PINK)))
    ws.conditional_formatting.add(f"Q5:Q{last}", CellIsRule(
        operator="equal", formula=['"Men pay more"'], font=Font(bold=True, color=BLUE)))
    ws.freeze_panes = "E5"
    widths(ws, [11, 10, 14, 20, 10, 10, 12, 12, 12, 12, 12, 12, 11, 10, 10, 11, 20])
    summary_rows = last

    # ---------------------------------------------------------------- Segment_Pivot
    ws = wb.create_sheet("Segment_Pivot")
    title(ws, "Pivot-style summary: segment x gender (adults)",
          "Built with COUNTIFS / AVERAGEIFS so it updates if the Data sheet changes.")
    header(ws, 4, ["Segment", "Products (Men)", "Products (Women)", "Avg MRP (Men)",
                   "Avg MRP (Women)", "Premium %", "Avg selling (Men)", "Avg selling (Women)",
                   "Premium after discount %", "Avg discount Men %", "Avg discount Women %"])
    segs = ["Apparel", "Footwear", "Accessories", "Personal Care"]
    SEG = rng("E")
    for i, s in enumerate(segs):
        row = 5 + i
        crit = f'{SEG},$A{row},{AGE},"Adult"'
        ws.append([s,
                   f'=COUNTIFS({crit},{GEN},"Men")', f'=COUNTIFS({crit},{GEN},"Women")',
                   f'=AVERAGEIFS({MRP},{crit},{GEN},"Men")', f'=AVERAGEIFS({MRP},{crit},{GEN},"Women")',
                   f"=(E{row}/D{row}-1)*100",
                   f'=AVERAGEIFS({SP},{crit},{GEN},"Men")', f'=AVERAGEIFS({SP},{crit},{GEN},"Women")',
                   f"=(H{row}/G{row}-1)*100",
                   f'=AVERAGEIFS({DISC},{crit},{GEN},"Men")', f'=AVERAGEIFS({DISC},{crit},{GEN},"Women")'])
    tr = 5 + len(segs)
    ws.append(["Total", f"=SUM(B5:B{tr - 1})", f"=SUM(C5:C{tr - 1})"])
    ws[f"A{tr}"].font = Font(bold=True)
    for col, fmt in {"D": RUPEE, "E": RUPEE, "G": RUPEE, "H": RUPEE, "F": PCT, "I": PCT,
                     "J": PCT, "K": PCT}.items():
        for cell in ws[col][4:tr]:
            cell.number_format = fmt
    widths(ws, [16, 12, 12, 13, 13, 11, 13, 13, 14, 12, 12])

    # ---------------------------------------------------------------- Matched_Pairs
    ws = wb.create_sheet("Matched_Pairs")
    title(ws, f"Brand-matched pairs ({len(pairs)}): same brand, same category",
          "Median MRP of the brand's women's line vs its men's line.")
    header(ws, 4, ["Source", "Age group", "Segment", "Category", "Brand", "Items (Men)",
                   "Items (Women)", "Median MRP (Men)", "Median MRP (Women)", "Gap %", "Outcome"])
    p = pairs.sort_values("ratio", ascending=False).reset_index(drop=True)
    for i, r in p.iterrows():
        row = 5 + i
        ws.append([r.source, r.age_group, r.segment, r.category, r.brand, int(r.n_men),
                   int(r.n_women), r.median_men, r.median_women, f"=(I{row}/H{row}-1)*100", r.outcome])
    lp = 4 + len(p)
    for col, fmt in {"H": RUPEE, "I": RUPEE, "J": PCT}.items():
        for cell in ws[col][4:lp]:
            cell.number_format = fmt
    ws.conditional_formatting.add(f"J5:J{lp}", DataBarRule(
        start_type="num", start_value=-60, end_type="num", end_value=60, color=PINK))
    ws.auto_filter.ref = f"A4:K{lp}"
    ws.freeze_panes = "A5"
    widths(ws, [11, 10, 14, 18, 22, 10, 10, 13, 13, 10, 22])

    # ---------------------------------------------------------------- Stats_Tests
    ws = wb.create_sheet("Stats_Tests")
    title(ws, "Statistical tests", "Fixed-effects OLS on log(MRP); SE clustered by brand.")
    header(ws, 4, ["Model", "Products", "Groups", "Brands", "Coefficient (Female)", "Cluster SE",
                   "Women's premium %", "95% CI low %", "95% CI high %", "p-value"])
    for r in reg.itertuples():
        ws.append([r.model, r.n_products, r.n_groups, r.n_brands, round(r.coef_female, 4),
                   round(r.se_cluster, 4), round(r.premium_pct, 2), round(r.ci_low_pct, 2),
                   round(r.ci_high_pct, 2), float(f"{r.p_value:.3g}")])
    start = 6 + len(reg)
    ws.cell(row=start, column=1, value="Brand-matched pair tests").font = Font(bold=True, size=12)
    header(ws, start + 1, ["Scope", "Pairs", "Women dearer %", "Men dearer %", "Within 2 %",
                           "Geo-mean premium %", "CI low %", "CI high %", "Wilcoxon p", "Sign-test p"])
    for r in psum.itertuples():
        ws.append([r.scope, r.pairs, round(r.women_dearer_pct, 1), round(r.men_dearer_pct, 1),
                   round(r.within_2pct_pct, 1), round(r.geo_mean_premium_pct, 2),
                   round(r.ci_low_pct, 2), round(r.ci_high_pct, 2),
                   float(f"{r.p_wilcoxon:.3g}"), float(f"{r.p_sign_test:.3g}")])
    widths(ws, [48, 11, 12, 12, 14, 12, 14, 12, 12, 12])

    # ---------------------------------------------------------------- Charts
    ws = wb.create_sheet("Charts")
    title(ws, "Charts")
    ch = BarChart()
    ch.type = "bar"
    ch.title = "Median premium % by category (women vs men)"
    ch.y_axis.title = "Premium %"
    ch.height, ch.width = 26, 22
    ch.x_axis.tickLblPos = "low"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.gapWidth = 40
    data = Reference(wb["Category_Summary"], min_col=13, min_row=4, max_row=summary_rows)
    labels = Reference(wb["Category_Summary"], min_col=4, min_row=5, max_row=summary_rows)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(labels)
    ch.series[0].graphicalProperties.solidFill = PINK
    ch.legend = None
    ws.add_chart(ch, "A3")

    ch2 = BarChart()
    ch2.title = "Average MRP by segment (adults)"
    ch2.height, ch2.width = 9, 16
    sp = wb["Segment_Pivot"]
    ch2.add_data(Reference(sp, min_col=4, max_col=5, min_row=4, max_row=8), titles_from_data=True)
    ch2.set_categories(Reference(sp, min_col=1, min_row=5, max_row=8))
    ch2.series[0].graphicalProperties.solidFill = BLUE
    ch2.series[1].graphicalProperties.solidFill = PINK
    ws.add_chart(ch2, "P3")

    ch3 = BarChart()
    ch3.title = "Average discount % by segment (adults)"
    ch3.height, ch3.width = 9, 16
    ch3.add_data(Reference(sp, min_col=10, max_col=11, min_row=4, max_row=8), titles_from_data=True)
    ch3.set_categories(Reference(sp, min_col=1, min_row=5, max_row=8))
    ch3.series[0].graphicalProperties.solidFill = BLUE
    ch3.series[1].graphicalProperties.solidFill = PINK
    ws.add_chart(ch3, "P22")

    out = EXCEL / "PinkTax_Analysis.xlsx"
    wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
